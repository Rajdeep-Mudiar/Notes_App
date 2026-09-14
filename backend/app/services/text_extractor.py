import io
import logging
import os
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)


class TextExtractor:
    """Multi-format text extraction engine for PDFs, DOCX, PPTX, plain text, and Note blocks."""

    @staticmethod
    def extract_from_file(file_path: str, mime_type: Optional[str] = None, original_name: Optional[str] = None) -> List[Dict[str, Any]]:
        """Extract text segments from a file given its local storage path.
        Returns list of dicts: [{'page_or_section': str, 'text': str}]
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at path: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        if original_name:
            ext = os.path.splitext(original_name)[1].lower() or ext

        if ext == ".pdf" or (mime_type and "pdf" in mime_type):
            return TextExtractor.extract_from_pdf(file_path)
        elif ext in [".docx", ".doc"] or (mime_type and "word" in mime_type):
            return TextExtractor.extract_from_docx(file_path)
        elif ext in [".pptx", ".ppt"] or (mime_type and ("presentation" in mime_type or "powerpoint" in mime_type)):
            return TextExtractor.extract_from_pptx(file_path)
        else:
            return TextExtractor.extract_from_plain_text(file_path)

    @staticmethod
    def extract_from_pdf(file_path: str) -> List[Dict[str, Any]]:
        """Extract text from PDF page by page."""
        sections: List[Dict[str, Any]] = []
        try:
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                clean_text = text.strip()
                if clean_text:
                    sections.append({
                        "page_or_section": f"Page {page_idx + 1}",
                        "text": clean_text
                    })
        except Exception as e:
            logger.warning(f"pypdf extraction failed for {file_path}: {e}. Attempting fallback plain text read.")
            try:
                with open(file_path, "rb") as f:
                    raw = f.read().decode("latin-1", errors="ignore")
                    # Filter printable text
                    printable = "".join(c for c in raw if 32 <= ord(c) <= 126 or c in "\n\r\t")
                    if len(printable.strip()) > 50:
                        sections.append({"page_or_section": "Page 1", "text": printable.strip()})
            except Exception as inner_e:
                logger.error(f"Fallback PDF text extraction failed: {inner_e}")

        if not sections:
            sections.append({"page_or_section": "Page 1", "text": "Document content without extractable text layer."})
        return sections

    @staticmethod
    def extract_from_docx(file_path: str) -> List[Dict[str, Any]]:
        """Extract text paragraphs and tables from DOCX."""
        sections: List[Dict[str, Any]] = []
        try:
            import docx

            doc = docx.Document(file_path)
            current_section = "Document Body"
            current_paras: List[str] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                # Check if paragraph is a heading
                if p.style and p.style.name.startswith("Heading"):
                    if current_paras:
                        sections.append({
                            "page_or_section": current_section,
                            "text": "\n".join(current_paras)
                        })
                        current_paras = []
                    current_section = text
                else:
                    current_paras.append(text)

            # Also extract tables
            for table_idx, table in enumerate(doc.tables):
                table_lines = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))
                if table_lines:
                    sections.append({
                        "page_or_section": f"Table {table_idx + 1}",
                        "text": "\n".join(table_lines)
                    })

            if current_paras:
                sections.append({
                    "page_or_section": current_section,
                    "text": "\n".join(current_paras)
                })
        except Exception as e:
            logger.warning(f"python-docx extraction failed for {file_path}: {e}")
            sections.append({"page_or_section": "Document Body", "text": "DOCX content extraction fallback."})

        if not sections:
            sections.append({"page_or_section": "Document Body", "text": "Document content is empty."})
        return sections

    @staticmethod
    def extract_from_pptx(file_path: str) -> List[Dict[str, Any]]:
        """Extract slide titles and shape text from PPTX."""
        sections: List[Dict[str, Any]] = []
        try:
            from pptx import Presentation

            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                slide_texts: List[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())

                if slide_texts:
                    sections.append({
                        "page_or_section": f"Slide {slide_idx + 1}",
                        "text": "\n".join(slide_texts)
                    })
        except Exception as e:
            logger.warning(f"python-pptx extraction failed for {file_path}: {e}")
            sections.append({"page_or_section": "Slide 1", "text": "PPTX content extraction fallback."})

        if not sections:
            sections.append({"page_or_section": "Slide 1", "text": "Slide presentation content is empty."})
        return sections

    @staticmethod
    def extract_from_plain_text(file_path: str) -> List[Dict[str, Any]]:
        """Extract plain text from UTF-8 / ASCII text files."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().strip()
                if content:
                    return [{"page_or_section": "Section 1", "text": content}]
        except Exception as e:
            logger.error(f"Plain text extraction failed for {file_path}: {e}")
        return [{"page_or_section": "Section 1", "text": "File content could not be decoded as text."}]

    @staticmethod
    def extract_from_note_blocks(title: str, blocks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Extract structured markdown text from Notion-style note blocks."""
        sections: List[Dict[str, Any]] = []
        current_section_title = f"Note: {title}"
        current_lines: List[str] = [f"# {title}"]

        for block in blocks:
            b_type = block.get("type", "paragraph")
            content = block.get("content", "").strip()
            props = block.get("properties", {}) or {}

            if not content and b_type not in ["divider", "table"]:
                continue

            if b_type in ["heading1", "heading_1"]:
                if len(current_lines) > 1:
                    sections.append({
                        "page_or_section": current_section_title,
                        "text": "\n".join(current_lines)
                    })
                    current_lines = []
                current_section_title = content
                current_lines.append(f"# {content}")
            elif b_type in ["heading2", "heading_2"]:
                current_lines.append(f"## {content}")
            elif b_type in ["heading3", "heading_3"]:
                current_lines.append(f"### {content}")
            elif b_type == "bullet":
                current_lines.append(f"* {content}")
            elif b_type == "numbered":
                current_lines.append(f"1. {content}")
            elif b_type == "checklist":
                checked = props.get("checked", False)
                box = "[x]" if checked else "[ ]"
                current_lines.append(f"- {box} {content}")
            elif b_type == "code":
                lang = props.get("language", "")
                current_lines.append(f"```{lang}\n{content}\n```")
            elif b_type == "quote":
                current_lines.append(f"> {content}")
            elif b_type == "callout":
                icon = props.get("icon", "💡")
                current_lines.append(f"> {icon} **Note:** {content}")
            elif b_type == "math":
                current_lines.append(f"$$\n{content}\n$$")
            elif b_type == "divider":
                current_lines.append("---")
            else:
                current_lines.append(content)

        if current_lines:
            sections.append({
                "page_or_section": current_section_title,
                "text": "\n".join(current_lines)
            })

        if not sections:
            sections.append({
                "page_or_section": f"Note: {title}",
                "text": f"# {title}\n(Empty note)"
            })
        return sections
